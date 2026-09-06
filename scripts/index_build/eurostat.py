"""Eurostat metropolitan regions: the pinned inputs and the offline derivation
of each city's metro from them.

The composition table (Eurostat's ``NUTS2021.xlsx``, whose ``Metropolitan``
sheet lists each NUTS-3 region with the metro code it belongs to, if any) and
the GISCO NUTS-3 polygons of the same NUTS revision are fetched once, checked
against pinned checksums and published together as the ``raw/eurostat.json``
generation; the readers parse the generation's verified bytes, so what was
hashed under this build's pins is what gets parsed. A city's metro is the one
whose composition holds the NUTS-3 region containing the representative point
of its Overture land areas.
"""

import collections
import io
import math
import re

import pyarrow as pa
import pyarrow.parquet as pq
import shapely
from shapely.strtree import STRtree

from index_build import geometry, pinned

NUTS_VERSION = "2021"
NUTS_SCALE = "01M"
POINTER = "eurostat.json"

COMPOSITION_FILE = "NUTS2021.xlsx"
COMPOSITION_URL = "https://ec.europa.eu/eurostat/documents/345175/629341/NUTS2021.xlsx"
COMPOSITION_SHEET = "Metropolitan"
COMPOSITION_HEADER = ("NUTS ID", "METRO (No/Yes)", "METRO CODE", "METRO LABEL")

BOUNDARIES_FILE = f"NUTS_RG_{NUTS_SCALE}_{NUTS_VERSION}_4326_LEVL_3.parquet"
BOUNDARIES_URL = (
    "https://gisco-services.ec.europa.eu/distribution/v2/nuts/parquet/"
    + BOUNDARIES_FILE
)
BOUNDARY_COLUMNS = ["NUTS_ID", "LEVL_CODE", "Shape"]

# The bytes verified live on 2026-09-06; a fetch that differs is refused.
PINS = {
    COMPOSITION_FILE: (
        "b17dcc379bb3586550ec3b16f3f474dd2a8a55b4cfdcdc1e1575097f1c2d4761"
    ),
    BOUNDARIES_FILE: (
        "0356e77f0903b101a03b29ddb76a6fcdb8314405048333beb8f2dcba0b6fb701"
    ),
}
URLS = {COMPOSITION_FILE: COMPOSITION_URL, BOUNDARIES_FILE: BOUNDARIES_URL}

# Country code, three digits, then M (metro) or MC (capital-city metro).
METRO_CODE = re.compile(r"\A([A-Z]{2})[0-9]{3}MC?\Z")
# Eurostat's code prefixes that are not the ISO 3166-1 alpha-2 codes the
# gazetteer uses; every other prefix is the country code itself.
EUROSTAT_COUNTRY = {"EL": "GR", "UK": "GB"}


class EurostatError(pinned.PinnedInputError):
    """A pinned Eurostat input is missing, altered or inconsistent."""


def prepare_inputs(cache_dir, *, files=None, expected=PINS):
    """Ensure ``raw/eurostat.json`` holds the pinned inputs; return its manifest.
    See :func:`pinned.prepare`."""
    return pinned.prepare(
        cache_dir,
        pointer=POINTER,
        urls=URLS,
        expected=expected,
        files=files,
        manifest={"source": "eurostat", "nuts_version": NUTS_VERSION},
        error=EurostatError,
    )


def resolve_inputs(cache_dir, *, expected=PINS):
    """``(generation, manifest)`` of the published inputs under this build's
    pins; see :func:`pinned.resolve`."""
    return pinned.resolve(
        cache_dir, pointer=POINTER, expected=expected, error=EurostatError
    )


# The composition's metros, and every NUTS-3 id its sheet lists (metro or
# not) — the set the boundary file must match exactly.
Composition = collections.namedtuple("Composition", ["metros", "nuts_ids"])


def _cell(value):
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def read_composition(data):
    """The workbook's metros and every NUTS-3 id its sheet lists.

    ``metros`` is ``{metro_code: {"name", "country", "nuts3": [...]}}`` with
    the country as the gazetteer's ISO code. The ``Metropolitan`` sheet is
    read as the input contract it is: the header, one row per NUTS-3 region,
    a code on every metro row and none on the others, the code's shape and
    one label per code. A reissued table that breaks any of these refuses
    the build rather than silently changing what a metro means.
    """
    import openpyxl

    try:
        workbook = openpyxl.load_workbook(
            io.BytesIO(data), read_only=True, data_only=True
        )
    except Exception as error:  # noqa: B902 - openpyxl raises its own hierarchy
        raise EurostatError(f"{COMPOSITION_FILE}: not a workbook: {error}") from None
    try:
        if COMPOSITION_SHEET not in workbook.sheetnames:
            raise EurostatError(f"{COMPOSITION_FILE}: no {COMPOSITION_SHEET!r} sheet")
        rows = workbook[COMPOSITION_SHEET].iter_rows(values_only=True)
        header = tuple(_cell(value) or "" for value in next(rows, ()))
        if header[: len(COMPOSITION_HEADER)] != COMPOSITION_HEADER:
            raise EurostatError(f"{COMPOSITION_FILE}: header {header!r}")
        metros = {}
        seen = set()
        for position, row in enumerate(rows, 2):
            cells = tuple(row) + (None,) * len(COMPOSITION_HEADER)
            nuts_id, flag, code, label = (_cell(value) for value in cells[:4])
            if nuts_id is None:
                continue
            if nuts_id in seen:
                raise EurostatError(
                    f"{COMPOSITION_FILE}: row {position}: NUTS-3 {nuts_id!r} twice"
                )
            seen.add(nuts_id)
            if flag == "N" and code is None:
                continue
            if flag != "Y" or code is None:
                raise EurostatError(
                    f"{COMPOSITION_FILE}: row {position}: flag {flag!r}, code {code!r}"
                )
            if label is None:
                raise EurostatError(
                    f"{COMPOSITION_FILE}: row {position}: {code} has no label"
                )
            match = METRO_CODE.match(code)
            if match is None:
                raise EurostatError(
                    f"{COMPOSITION_FILE}: row {position}: metro code {code!r}"
                )
            prefix = match.group(1)
            metro = metros.setdefault(
                code,
                {
                    "name": label,
                    "country": EUROSTAT_COUNTRY.get(prefix, prefix),
                    "nuts3": [],
                },
            )
            if label != metro["name"]:
                raise EurostatError(
                    f"{COMPOSITION_FILE}: row {position}: {code} is labelled "
                    f"{label!r} and {metro['name']!r}"
                )
            metro["nuts3"].append(nuts_id)
    finally:
        workbook.close()
    if not metros:
        raise EurostatError(f"{COMPOSITION_FILE}: no metropolitan regions")
    for metro in metros.values():
        metro["nuts3"].sort()
    return Composition(metros, sorted(seen))


def read_boundaries(data):
    """``{nuts_id: geometry}`` for the level-3 regions of the GISCO file."""
    try:
        table = pq.read_table(pa.BufferReader(data), columns=BOUNDARY_COLUMNS)
    except Exception as error:  # noqa: B902 - pyarrow raises its own hierarchy
        raise EurostatError(f"{BOUNDARIES_FILE}: not readable: {error}") from None
    boundaries = {}
    for row in table.to_pylist():
        nuts_id = row["NUTS_ID"]
        if row["LEVL_CODE"] != 3 or not nuts_id:
            raise EurostatError(f"{BOUNDARIES_FILE}: {nuts_id!r} is not a NUTS-3 row")
        if nuts_id in boundaries:
            raise EurostatError(f"{BOUNDARIES_FILE}: {nuts_id!r} twice")
        try:
            geom = shapely.force_2d(shapely.from_wkb(row["Shape"]))
        except Exception:  # noqa: B902 - shapely raises its own hierarchy
            raise EurostatError(f"{BOUNDARIES_FILE}: {nuts_id!r} geometry") from None
        if not geometry._valid_polygon(geom):
            raise EurostatError(
                f"{BOUNDARIES_FILE}: {nuts_id!r} is not a valid polygon"
            )
        boundaries[nuts_id] = geom
    if not boundaries:
        raise EurostatError(f"{BOUNDARIES_FILE}: no regions")
    return boundaries


def load_inputs(cache_dir, *, expected=PINS):
    """``(metros, boundaries, manifest)`` parsed from the verified bytes of
    the generation :func:`resolve_inputs` accepts. The sheet's NUTS-3 ids and
    the boundary file's must match both ways: a revision drift between the
    two pins refuses rather than remaps."""
    generation, manifest = resolve_inputs(cache_dir, expected=expected)
    with generation:
        composition = read_composition(generation.read_bytes(COMPOSITION_FILE))
        boundaries = read_boundaries(generation.read_bytes(BOUNDARIES_FILE))
    sheet, drawn = set(composition.nuts_ids), set(boundaries)
    missing, unexpected = sorted(sheet - drawn), sorted(drawn - sheet)
    if missing or unexpected:
        raise EurostatError(
            f"composition and NUTS {NUTS_VERSION} boundaries disagree: "
            f"{len(missing)} sheet codes without a boundary {missing[:10]}, "
            f"{len(unexpected)} boundary ids not in the sheet {unexpected[:10]}"
        )
    return composition.metros, boundaries, manifest


def countries(metros):
    """The countries the composition covers, as the gazetteer's codes."""
    return {metro["country"] for metro in metros.values()}


class Containment:
    """Which NUTS-3 regions a point falls in, over an STRtree of the boundaries."""

    def __init__(self, boundaries):
        self._ids = sorted(boundaries)
        self._tree = STRtree([boundaries[nuts_id] for nuts_id in self._ids])

    def regions_at(self, point):
        """The NUTS-3 ids covering the point, sorted; several on a boundary."""
        found = self._tree.query(point, predicate="covered_by")
        return sorted(self._ids[index] for index in found)


def _footprint(rows):
    """The union of a place's usable land areas, or None."""
    geoms = [
        row["geom"]
        for row in rows or []
        if row["geom"] is not None and not row["geom"].is_empty
    ]
    if not geoms:
        return None
    merged = geoms[0] if len(geoms) == 1 else shapely.unary_union(geoms)
    return None if merged.is_empty else merged


def _pick(candidates, footprint, boundaries, by_nuts3):
    """``(nuts_id, ambiguous)`` among the regions covering a footprint's
    representative point. One candidate, or several implying one metro,
    settle it; otherwise the region holding the larger share of the
    footprint wins, and a tie (or a footprint without area) is ambiguous —
    a boundary point may not choose between metros on its own."""
    if len(candidates) <= 1:
        return (candidates[0] if candidates else None), False
    if len({by_nuts3.get(nuts_id) for nuts_id in candidates}) == 1:
        return candidates[0], False
    shares = sorted(
        (
            (footprint.intersection(boundaries[nuts_id]).area, nuts_id)
            for nuts_id in candidates
        ),
        reverse=True,
    )
    (top, best), (second, _) = shares[0], shares[1]
    if top == 0 or math.isclose(top, second, rel_tol=1e-9):
        return None, True
    return best, False


def assign(places, areas, metros, boundaries):
    """One assignment row per city of a covered country, by ``city_id``.

    ``status`` is ``"assigned"`` with the ``metro_code`` whose composition
    holds the city's NUTS-3 region, ``"unassigned"`` when that region is in
    no metro (or the city lies outside every region), ``"ambiguous"`` when
    the regions covering its representative point imply different metros
    and hold equal shares of its footprint, and ``"unplaceable"`` when the
    city has no usable land area. ``areas`` is ``geometry.read_areas``
    output keyed by Overture id.
    """
    by_nuts3 = {
        nuts_id: code for code, metro in metros.items() for nuts_id in metro["nuts3"]
    }
    covered = countries(metros)
    containment = Containment(boundaries)
    rows = []
    for place in places:
        if place.get("kind") != "city" or place.get("country_code") not in covered:
            continue
        overture_id = place.get("overture_id")
        footprint = _footprint(areas.get(overture_id)) if overture_id else None
        if footprint is None:
            status, nuts_id = "unplaceable", None
        else:
            candidates = containment.regions_at(footprint.representative_point())
            nuts_id, ambiguous = _pick(candidates, footprint, boundaries, by_nuts3)
            if ambiguous:
                status = "ambiguous"
            elif by_nuts3.get(nuts_id):
                status = "assigned"
            else:
                status = "unassigned"
        rows.append(
            {
                "city_id": place["place_id"],
                "status": status,
                "nuts_id": nuts_id,
                "metro_code": by_nuts3.get(nuts_id),
            }
        )
    return sorted(rows, key=lambda row: row["city_id"])
